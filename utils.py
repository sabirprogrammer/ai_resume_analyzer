"""Document extraction and normalization. No UI or network dependencies."""
import io
import re
import unicodedata
import zipfile
from pathlib import Path

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARS = 30000
MAX_JOB_CHARS = 20000


class InputError(ValueError):
    """An input the user can correct."""


def clean_text(text: str) -> str:
    text = unicodedata.normalize('NFKC', text).replace('\x00', '')
    return '\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines()).strip()


def normalize_skill(value: str) -> str:
    value = re.sub(r'\s+', ' ', clean_text(value).casefold()).strip()
    return {'js': 'javascript', 'ts': 'typescript', 'nodejs': 'node.js',
            'node js': 'node.js', 'reactjs': 'react', 'react.js': 'react',
            'postgres': 'postgresql', 'amazon web services': 'aws',
            'rest apis': 'rest api'}.get(value, value)


def contains_phrase(phrase: str, text: str) -> bool:
    """Token boundaries prevent Java matching JavaScript, or C matching C++."""
    pattern = r'(?<![\w+#])' + re.escape(clean_text(phrase).casefold()).replace(r'\ ', r'\s+') + r'(?![\w+#])'
    return bool(phrase.strip()) and bool(re.search(pattern, clean_text(text).casefold()))


def extract_pdf_text(data: bytes) -> str:
    import fitz
    with fitz.open(stream=data, filetype='pdf') as document:
        if document.needs_pass:
            raise InputError('This PDF is password protected. Upload an unlocked copy.')
        if len(document) > 20:
            raise InputError('Please upload a resume of 20 pages or fewer.')
        pages = [page.get_text(sort=True) for page in document]
        if any(len(page.strip()) < 20 for page in pages):
            raise InputError('A PDF page has little or no readable text. Use a text-based PDF or DOCX; scanned pages need OCR first.')
        return '\n'.join(pages)


def extract_docx_text(data: bytes) -> str:
    from docx import Document
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
            raise InputError('The expanded DOCX is too large.')
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts.extend(' | '.join(cell.text for cell in row.cells) for row in table.rows)
    for section in document.sections:
        for area in (section.header, section.footer):
            parts.extend(p.text for p in area.paragraphs)
    return '\n'.join(parts)


def extract_resume_text(data: bytes, filename: str) -> str:
    if not data or len(data) > MAX_FILE_BYTES:
        raise InputError('Upload a non-empty file smaller than 5 MB.')
    suffix = Path(filename).suffix.lower()
    readers = {'.pdf': extract_pdf_text, '.docx': extract_docx_text}
    if suffix not in readers:
        raise InputError('Only PDF and DOCX resumes are supported.')
    try:
        text = clean_text(readers[suffix](data))
    except InputError:
        raise
    except Exception as exc:
        raise InputError('Unable to read this file. Export a fresh PDF or DOCX and try again.') from exc
    if len(text) < 100 or len(text.split()) < 20:
        raise InputError('Not enough readable resume text. Scanned files need OCR first.')
    if len(text) > MAX_RESUME_CHARS:
        raise InputError('Resume text exceeds 30,000 characters. Upload a shorter version.')
    return text


def validate_inputs(resume: str, job: str) -> tuple[str, str]:
    resume, job = clean_text(resume), clean_text(job)
    for name, value, maximum in [('Resume', resume, MAX_RESUME_CHARS), ('Job description', job, MAX_JOB_CHARS)]:
        if len(value) < 100 or len(value.split()) < 20:
            raise InputError(f'{name} needs at least 100 characters and 20 words.')
        if len(value) > maximum:
            raise InputError(f'{name} exceeds {maximum:,} characters.')
    return resume, job
