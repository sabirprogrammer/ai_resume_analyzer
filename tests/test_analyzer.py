import io
import hashlib
from pathlib import Path
from unittest.mock import Mock, patch
import fitz
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest
from analyzer import (Advice, AnalysisError, DEFAULT_MODEL, Extraction, SenseNovaClient, analyze_resume,
                      calculate_score, compare_skills, validate_evidence)
from utils import InputError, contains_phrase, extract_resume_text

RESUME = '''Alex Example alex@example.com
Skills: Python, SQL, JavaScript
Experience: Built 12 Python applications over 2 years.
Projects: Created a reporting system with SQL and tested multiple software components.
Education: Bachelor of Computer Science.'''
JOB = '''We require Python and Docker for a software developer role. Applicants need 2 years of development experience.
A Bachelor of Computer Science is required. You will build reliable systems and work with our engineering team.'''


def extraction():
    return Extraction.model_validate({
        'resume_skills': [{'name': 'Python', 'quote': 'Skills: Python, SQL, JavaScript'}],
        'required_skills': [{'name': 'Python', 'quote': 'We require Python and Docker'},
                            {'name': 'Docker', 'quote': 'We require Python and Docker'}],
        'preferred_skills': [],
        'ats_keywords': [{'name': 'Python', 'quote': 'We require Python and Docker'},
                         {'name': 'Docker', 'quote': 'We require Python and Docker'}],
        'experience': [{'requirement': '2 years', 'job_quote': '2 years of development experience',
                        'status': 'met', 'resume_quote': 'Built 12 Python applications over 2 years.', 'explanation': 'Two years stated.'}],
        'education': [{'requirement': 'Bachelor of Computer Science', 'job_quote': 'A Bachelor of Computer Science is required.',
                       'status': 'met', 'resume_quote': 'Bachelor of Computer Science', 'explanation': 'Degree matches.'}]})


def test_score_and_evidence():
    e = validate_evidence(extraction(), RESUME, JOB)
    comparison = compare_skills(e, RESUME)
    assert comparison['missing_skills'] == ['Docker']
    assert calculate_score(e, comparison, RESUME)['match_score'] == 70


def test_missing_categories_renormalized():
    e = extraction()
    e.experience, e.education, e.ats_keywords = [], [], []
    result = calculate_score(e, compare_skills(e, RESUME), RESUME)
    assert result['match_score'] == 61.1
    assert result['score_breakdown']['education']['rate'] is None


def test_boundaries():
    assert not contains_phrase('Java', 'JavaScript')
    assert not contains_phrase('C', 'C++')
    assert contains_phrase('C++', 'Uses C++ and Python')
    assert contains_phrase('SQL', 'SQL, Python')


def test_hallucination_rejected():
    e = extraction()
    e.resume_skills[0].quote = 'I have Kubernetes experience'
    with pytest.raises(AnalysisError):
        validate_evidence(e, RESUME, JOB)


def test_duplicate_normalization():
    e = extraction()
    e.required_skills.append(e.required_skills[0])
    assert len(validate_evidence(e, RESUME, JOB).required_skills) == 2


def test_docx_extraction():
    document = Document()
    document.add_paragraph(RESUME)
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Table skill: Docker'
    buffer = io.BytesIO()
    document.save(buffer)
    assert 'Table skill: Docker' in extract_resume_text(buffer.getvalue(), 'resume.docx')


def test_pdf_extraction_and_scan():
    document = fitz.open()
    document.new_page().insert_text((40, 40), RESUME)
    assert 'Python' in extract_resume_text(document.tobytes(), 'resume.pdf')
    blank = fitz.open()
    blank.new_page()
    with pytest.raises(InputError, match='readable text'):
        extract_resume_text(blank.tobytes(), 'scan.pdf')


@pytest.mark.parametrize('data,name', [(b'', 'resume.pdf'), (b'bad', 'resume.docx'), (b'bad', 'resume.txt')])
def test_invalid_uploads(data, name):
    with pytest.raises(InputError):
        extract_resume_text(data, name)


def response(content, status=200, finish='stop'):
    return Mock(status_code=status, json=lambda: {'choices': [{'message': {'content': content}, 'finish_reason': finish}]})


def test_json_regeneration():
    with patch('analyzer.requests.post', side_effect=[response('invalid'), response(extraction().model_dump_json())]) as post:
        assert SenseNovaClient('fake').generate('extract', {}, Extraction)
        assert post.call_count == 2
        assert post.call_args.args[0] == 'https://token.sensenova.ai/v1/chat/completions'


def test_auth_error_sanitized():
    with patch('analyzer.requests.post', return_value=response('secret', 401)):
        with pytest.raises(AnalysisError, match='HTTP 401'):
            SenseNovaClient('fake').generate('extract', {}, Extraction)


def test_http_backoff():
    with patch('analyzer.requests.post', side_effect=[response('', 429), response(extraction().model_dump_json())]), patch('analyzer.time.sleep'):
        assert SenseNovaClient('fake').generate('extract', {}, Extraction)


def test_truncated_response():
    with patch('analyzer.requests.post', return_value=response('{}', finish='length')):
        with pytest.raises(AnalysisError, match='cut off'):
            SenseNovaClient('fake').generate('extract', {}, Extraction)


def test_complete_flow_and_advice_failure():
    client = Mock()
    client.generate.side_effect = [extraction(), Advice(strengths=['Python'], problems=['Docker not evidenced'], recommendations=['Learn Docker if new to you.'], final_assessment='Review the required Docker gap.')]
    result = analyze_resume(client, RESUME, JOB)
    assert result['match_score'] == 70
    assert result['recommendations']
    client.generate.side_effect = [extraction(), AnalysisError('Temporary failure')]
    assert analyze_resume(client, RESUME, JOB)['match_score'] == 70


def test_streamlit_initial_and_validation():
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py')).run()
    assert not app.exception
    app.button[0].click().run()
    assert not app.exception
    assert 'Upload your resume' in app.error[0].value


def test_streamlit_results():
    e = extraction()
    result = compare_skills(e, RESUME)
    result.update(calculate_score(e, result, RESUME))
    result.update(evidence=e.model_dump(), advice_warning='', strengths=['Python'], problems=[], recommendations=['Review Docker'], final_assessment='Good Match')
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'))
    app.session_state['fingerprint'] = hashlib.sha256(DEFAULT_MODEL.encode()).hexdigest()
    app.session_state['result'] = result
    app.session_state['resume_text'] = RESUME
    app.run()
    assert not app.exception
    assert app.metric[0].value == '70%'
